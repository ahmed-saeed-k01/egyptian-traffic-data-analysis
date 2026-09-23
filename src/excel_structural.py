"""
src/excel_structural.py

Traffic_Data - Excel Structural Reader
======================================

Purpose
-------
Read-only structural reader for the Excel sources accepted by Discovery.

This module sits AFTER discovery.py and BEFORE structural_detectors.py.

Architecture
------------

    config.py
        ↓
    discovery.py
        ↓
    file_manifest.csv
        ↓
    ExcelStructuralInput
        ↓
    excel_structural.py
        ↓
    Raw Excel Structural Evidence
        ↓
    structural_detectors.py
        ↓
    structural_scan.py


Primary responsibilities
------------------------
1. Consume Excel files selected by Discovery.
2. Validate the Discovery input contract.
3. Open .xlsx workbooks read-only.
4. Preserve formulas by using data_only=False.
5. Enumerate every worksheet in workbook order.
6. Preserve raw non-empty cell values and their positions.
7. Preserve Excel value types.
8. Record structural empty-cell spans without materializing empty cells.
9. Record row-level structural evidence.
10. Record worksheet dimensions and observed data boundaries.
11. Record merged cells.
12. Record hidden sheets, rows, and columns.
13. Record freeze panes.
14. Record workbook metadata when available.
15. Record preview information.
16. Record read errors without modifying the source.
17. Preserve Discovery provenance through file_id and relative_path.
18. Produce deterministic, JSON-serializable structural evidence.
19. Never perform semantic classification.

Non-responsibilities
--------------------
This module MUST NOT:

- choose the final header row
- rename columns
- clean values
- forward-fill values
- create analytical records
- remove totals
- remove notes
- identify Brand semantically
- compare headers with mapping.xlsx
- infer schema
- infer grain
- infer period
- modify raw Excel files
- overwrite raw files
- silently skip unreadable worksheets
- rescan data/raw as its primary input mechanism

Important design rule
---------------------
An empty Excel cell is not materialized as an individual CellEvidence object.

Instead:

    non-empty cells → stored individually
    empty cells     → represented structurally as spans

This is important for large workbooks.

The reader preserves evidence.
The detector interprets evidence.

Known openpyxl read-only limitations (fixed in this version)
--------------------------------------------------------------
openpyxl's read-only worksheet object (``ReadOnlyWorksheet``) does NOT
inherit from the normal ``Worksheet`` class. It only carries over a
handful of methods (``iter_rows``, ``cell``, ``values``, ``rows``).
It does not expose ``row_dimensions``, ``column_dimensions``,
``merged_cells``, or ``freeze_panes`` at all — accessing any of these
on a read-only worksheet raises ``AttributeError``. This is verified
against openpyxl 3.1.5 in this project's environment.

Because responsibilities 11-13 above (merged cells, hidden rows and
columns, freeze panes) are explicit requirements of this reader, and
dropping them silently would lose evidence, this module opens a
*second*, normal-mode handle to the same file solely to read that
worksheet-level metadata, and always keeps the primary read-only
handle for the memory-efficient row-by-row cell scan. If the second,
normal-mode open fails (e.g. on a very large or malformed file), the
sheet result is still returned with ``metadata_read_status`` set to
an explanatory value instead of crashing the file. This keeps a
"we didn't check" distinct from "we checked, there is none" evidence
state, which is what this reader is meant to preserve.

Additionally, ``worksheet.calculate_dimension()`` and
``worksheet.max_row`` / ``worksheet.max_column`` are only populated
from the workbook's optional ``<dimension>`` XML tag. Real-world
exports from non-Excel report generators (a realistic case for
traffic-authority monthly exports) can omit or mis-write this tag,
in which case these become ``None`` and ``calculate_dimension()``
raises ``ValueError``. Both are handled defensively below so that a
missing/incorrect dimension tag on one file degrades gracefully
instead of failing the whole file.
"""


from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable
import hashlib
import json
import logging
import math
import re
from numbers import Number

import openpyxl


# ============================================================================
# Optional project configuration
# ============================================================================

try:
    from src.config import (
        RAW_DIR,
        PREVIEW_ROWS,
        PIPELINE_VERSION,
    )
except ImportError:
    # The fallback exists only to make this module easier to import in
    # isolated development/testing environments.
    RAW_DIR = Path("data/raw")
    PREVIEW_ROWS = 15
    PIPELINE_VERSION = "unknown"


LOGGER = logging.getLogger(__name__)


# ============================================================================
# Constants
# ============================================================================

SUPPORTED_EXCEL_EXTENSION = ".xlsx"

READER_NAME = "openpyxl"
READER_VERSION = getattr(openpyxl, "__version__", "unknown")

# Bumped from 1.0 -> 1.1: ExcelSheetStructuralResult gained
# `metadata_read_status` (see module docstring, "Known openpyxl
# read-only limitations"). Any consumer keying off this field should
# check reader_schema_version.
READER_SCHEMA_VERSION = "1.1"

# Excel error literals.
EXCEL_ERROR_VALUES = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#GETTING_DATA",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#BLOCKED!",
    "#UNKNOWN!",
}

# XML-compatible control characters are not useful in JSON/text evidence.
_CONTROL_CHARS_RE = re.compile(
    r"[\x00-\x08\x0B\x0C\x0E-\x1F]"
)


# ============================================================================
# Input Contract
# ============================================================================

@dataclass(frozen=True)
class ExcelStructuralInput:
    """
    Minimal input contract between discovery.py and excel_structural.py.

    Discovery remains the source of truth.

    Only the metadata required to locate and identify the source file is
    accepted here. The complete pandas Manifest row is deliberately not
    passed into the reader.
    """

    file_id: str
    relative_path: str
    file_name: str
    extension: str
    source_root: str

    # These two fields are validation/provenance helpers.
    # They do not make the reader dependent on the complete Manifest schema.
    discovery_status: str = "DISCOVERED"
    file_type: str = "excel"

    def validate(self) -> None:
        """Validate the structural-reader input contract."""

        errors: list[str] = []

        if not self.file_id.strip():
            errors.append("file_id is required.")

        if not self.relative_path.strip():
            errors.append("relative_path is required.")

        if not self.file_name.strip():
            errors.append("file_name is required.")

        extension = self.extension.lower().strip()

        if extension != SUPPORTED_EXCEL_EXTENSION:
            errors.append(
                "Unsupported Excel extension: "
                f"{self.extension!r}. "
                "This project currently accepts .xlsx only."
            )

        if self.discovery_status != "DISCOVERED":
            errors.append(
                "Excel Structural accepts only Discovery records with "
                "discovery_status='DISCOVERED'."
            )

        if self.file_type.lower().strip() != "excel":
            errors.append(
                "Excel Structural accepts only file_type='excel'."
            )

        if errors:
            raise ValueError(
                "Invalid ExcelStructuralInput: "
                + " ".join(errors)
            )

    def resolve_path(self, project_root: Path | None = None) -> Path:
        """
        Resolve the physical source path from:

            source_root + relative_path

        No raw-directory scan is performed.
        """

        self.validate()

        root = Path(self.source_root)

        if not root.is_absolute():
            if project_root is None:
                project_root = Path(__file__).resolve().parents[1]

            root = project_root / root

        path = (root / self.relative_path).resolve()

        return path


# ============================================================================
# Raw Cell Evidence
# ============================================================================

@dataclass(frozen=True)
class ExcelCellEvidence:
    """
    Raw structural evidence for one non-empty Excel cell.

    No semantic interpretation is performed.
    """

    row_index: int
    column_index: int
    column_letter: str
    cell_address: str

    raw_value: Any
    value_type: str

    is_formula: bool = False
    formula: str | None = None

    data_type: str | None = None


# ============================================================================
# Empty Cell Span
# ============================================================================

@dataclass(frozen=True)
class ExcelEmptyCellSpan:
    """
    Structural representation of consecutive empty cells.

    Empty cells are represented as spans rather than individual objects.

    Example
    -------
    If B10:D10 are empty:

        row_index = 10
        start_column_index = 2
        end_column_index = 4

    This preserves the structural fact without creating three cell objects.
    """

    row_index: int
    start_column_index: int
    end_column_index: int
    start_column_letter: str
    end_column_letter: str
    length: int


# ============================================================================
# Row Evidence
# ============================================================================

@dataclass(frozen=True)
class ExcelRowEvidence:
    """
    Structural evidence for one observed worksheet row.
    """

    row_index: int
    excel_row: int

    non_empty_count: int
    empty_cell_count: int

    first_non_empty_column: int | None
    last_non_empty_column: int | None

    first_non_empty_column_letter: str | None
    last_non_empty_column_letter: str | None

    is_completely_empty: bool

    non_empty_cells: tuple[ExcelCellEvidence, ...] = ()

    empty_spans: tuple[ExcelEmptyCellSpan, ...] = ()


# ============================================================================
# Sheet Boundaries
# ============================================================================

@dataclass(frozen=True)
class ExcelSheetBoundaries:
    """
    Observed structural boundaries of a worksheet.

    Reported worksheet dimensions and observed data boundaries are deliberately
    separated because Excel's reported dimensions may exceed actual data.
    """

    reported_max_row: int
    reported_max_column: int

    observed_first_non_empty_row: int | None
    observed_last_non_empty_row: int | None

    observed_first_non_empty_column: int | None
    observed_last_non_empty_column: int | None

    observed_first_non_empty_column_letter: str | None
    observed_last_non_empty_column_letter: str | None

    observed_non_empty_row_count: int
    observed_non_empty_cell_count: int


# ============================================================================
# Workbook / Worksheet Metadata
# ============================================================================

@dataclass(frozen=True)
class ExcelMergedRangeEvidence:
    """One merged-cell range exactly as reported by Excel."""

    range_reference: str


@dataclass(frozen=True)
class ExcelHiddenRowEvidence:
    """One hidden worksheet row."""

    row_index: int


@dataclass(frozen=True)
class ExcelHiddenColumnEvidence:
    """One hidden worksheet column."""

    column_index: int
    column_letter: str


# ============================================================================
# Sheet Structural Result
# ============================================================================

@dataclass
class ExcelSheetStructuralResult:
    """
    Complete structural evidence for one worksheet.

    This is intentionally evidence-oriented.
    """

    sheet_name: str
    sheet_index: int
    sheet_state: str

    read_status: str

    reported_dimensions: str | None

    boundaries: ExcelSheetBoundaries | None

    rows: list[ExcelRowEvidence] = field(
        default_factory=list
    )

    empty_row_indices: list[int] = field(
        default_factory=list
    )

    empty_column_indices: list[int] = field(
        default_factory=list
    )

    empty_column_letters: list[str] = field(
        default_factory=list
    )

    merged_ranges: list[ExcelMergedRangeEvidence] = field(
        default_factory=list
    )

    hidden_rows: list[ExcelHiddenRowEvidence] = field(
        default_factory=list
    )

    hidden_columns: list[ExcelHiddenColumnEvidence] = field(
        default_factory=list
    )

    freeze_panes: str | None = None

    # NEW (schema 1.1): whether merged_ranges/hidden_rows/hidden_columns/
    # freeze_panes were actually inspected.
    #   "READ_OK"     -> the normal-mode metadata handle opened fine; an
    #                    empty list above means "genuinely no merged
    #                    cells / hidden rows / etc.", not "unchecked".
    #   "UNAVAILABLE: <error>" -> the second, normal-mode open failed
    #                    (e.g. very large or malformed file); the four
    #                    fields above are empty/None because they were
    #                    never inspected, not because the sheet has none.
    metadata_read_status: str = "UNAVAILABLE"

    preview: list[list[Any]] = field(
        default_factory=list
    )

    preview_rows_requested: int = PREVIEW_ROWS

    preview_rows_read: int = 0

    preview_columns: int = 0

    error: str | None = None


# ============================================================================
# Workbook Metadata
# ============================================================================

@dataclass(frozen=True)
class ExcelWorkbookMetadata:
    """
    Optional workbook properties.

    Metadata is informational only and never used to infer analytical meaning.
    """

    creator: str | None = None
    last_modified_by: str | None = None
    title: str | None = None
    subject: str | None = None
    description: str | None = None
    category: str | None = None
    keywords: str | None = None
    created: str | None = None
    modified: str | None = None


# ============================================================================
# Workbook Structural Result
# ============================================================================

@dataclass
class ExcelWorkbookStructuralResult:
    """
    Complete structural evidence for one discovered Excel workbook.
    """

    file_id: str
    file_name: str
    relative_path: str
    source_root: str
    extension: str

    read_status: str

    reader: str = READER_NAME
    reader_version: str = READER_VERSION
    reader_schema_version: str = READER_SCHEMA_VERSION
    pipeline_version: str = PIPELINE_VERSION

    sheet_count: int = 0

    sheets: list[ExcelSheetStructuralResult] = field(
        default_factory=list
    )

    workbook_metadata: ExcelWorkbookMetadata | None = None

    error: str | None = None


# ============================================================================
# Generic Serialization Helpers
# ============================================================================

def _safe_string(value: Any) -> str:
    """
    Convert a value to a deterministic human-readable string without
    changing the original value stored in raw evidence.
    """

    if value is None:
        return ""

    text = str(value)

    return _CONTROL_CHARS_RE.sub("", text)


def _json_safe_value(value: Any) -> Any:
    """
    Convert Excel/Python values into JSON-safe values.

    This is serialization only.

    It does NOT normalize analytical values.
    """

    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"

        if math.isinf(value):
            return "Infinity" if value > 0 else "-Infinity"

        return value

    if isinstance(value, Decimal):
        return str(value)

    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, time):
        return value.isoformat()

    if isinstance(value, bytes):
        return value.hex()

    if isinstance(value, str):
        return _CONTROL_CHARS_RE.sub("", value)

    if isinstance(value, (list, tuple)):
        return [
            _json_safe_value(item)
            for item in value
        ]

    if isinstance(value, dict):
        return {
            str(key): _json_safe_value(item)
            for key, item in value.items()
        }

    if isinstance(value, Number):
        return float(value)

    return _safe_string(value)


def _make_json_serializable(value: Any) -> Any:
    """
    Recursively convert dataclasses and values into JSON-safe objects.
    """

    if hasattr(value, "__dataclass_fields__"):
        return {
            key: _make_json_serializable(item)
            for key, item in asdict(value).items()
        }

    if isinstance(value, dict):
        return {
            str(key): _make_json_serializable(item)
            for key, item in value.items()
        }

    if isinstance(value, (list, tuple)):
        return [
            _make_json_serializable(item)
            for item in value
        ]

    return _json_safe_value(value)


# ============================================================================
# Excel Column Helpers
# ============================================================================

def excel_column_letter(column_index: int) -> str:
    """
    Convert one-based Excel column number to Excel letters.

    1 -> A
    26 -> Z
    27 -> AA
    """

    if column_index < 1:
        raise ValueError(
            "Excel column index must be >= 1."
        )

    result = ""

    value = column_index

    while value:
        value, remainder = divmod(
            value - 1,
            26,
        )
        result = chr(65 + remainder) + result

    return result


# ============================================================================
# Value Classification
# ============================================================================

def classify_excel_value(
    value: Any,
    cell_data_type: str | None,
) -> str:
    """
    Classify the raw Python/Excel value for structural evidence.

    No analytical conversion is performed.
    """

    if value is None:
        return "empty"

    if isinstance(value, bool):
        return "boolean"

    if isinstance(value, datetime):
        return "datetime"

    if isinstance(value, date):
        return "date"

    if isinstance(value, time):
        return "time"

    if isinstance(value, int):
        return "integer"

    if isinstance(value, float):
        return "float"

    if isinstance(value, Decimal):
        return "decimal"

    if isinstance(value, str):
        if value in EXCEL_ERROR_VALUES:
            return "excel_error"

        if value.startswith("="):
            return "formula"

        return "text"

    if cell_data_type == "e":
        return "excel_error"

    return type(value).__name__


# ============================================================================
# Cell Evidence
# ============================================================================

def _build_cell_evidence(cell: Any) -> ExcelCellEvidence:
    """
    Build raw structural evidence for one non-empty cell.
    """

    raw_value = cell.value

    value_type = classify_excel_value(
        raw_value,
        getattr(cell, "data_type", None),
    )

    is_formula = (
        getattr(cell, "data_type", None) == "f"
        or (
            isinstance(raw_value, str)
            and raw_value.startswith("=")
        )
    )

    formula = (
        raw_value
        if is_formula and isinstance(raw_value, str)
        else None
    )

    return ExcelCellEvidence(
        row_index=cell.row,
        column_index=cell.column,
        column_letter=excel_column_letter(cell.column),
        cell_address=cell.coordinate,
        raw_value=_json_safe_value(raw_value),
        value_type=value_type,
        is_formula=is_formula,
        formula=formula,
        data_type=getattr(
            cell,
            "data_type",
            None,
        ),
    )


# ============================================================================
# Empty Span Detection
# ============================================================================

def _build_empty_spans(
    row_number: int,
    non_empty_columns: set[int],
    max_column: int,
) -> tuple[ExcelEmptyCellSpan, ...]:
    """
    Represent empty cells in one row as consecutive column spans.

    Only spans inside the observed/reportable row width are represented.
    """

    if max_column <= 0:
        return ()

    spans: list[ExcelEmptyCellSpan] = []

    start: int | None = None

    for column_index in range(
        1,
        max_column + 1,
    ):
        is_empty = column_index not in non_empty_columns

        if is_empty and start is None:
            start = column_index

        elif not is_empty and start is not None:
            end = column_index - 1

            spans.append(
                ExcelEmptyCellSpan(
                    row_index=row_number,
                    start_column_index=start,
                    end_column_index=end,
                    start_column_letter=(
                        excel_column_letter(start)
                    ),
                    end_column_letter=(
                        excel_column_letter(end)
                    ),
                    length=end - start + 1,
                )
            )

            start = None

    if start is not None:
        end = max_column

        spans.append(
            ExcelEmptyCellSpan(
                row_index=row_number,
                start_column_index=start,
                end_column_index=end,
                start_column_letter=(
                    excel_column_letter(start)
                ),
                end_column_letter=(
                    excel_column_letter(end)
                ),
                length=end - start + 1,
            )
        )

    return tuple(spans)


# ============================================================================
# Preview
# ============================================================================

def _build_preview(
    worksheet: Any,
    preview_rows: int,
) -> list[list[Any]]:
    """
    Build a JSON-safe preview.

    Preview is supplemental evidence only.

    FIX: bound the read directly through ``iter_rows(max_row=preview_rows)``
    instead of computing ``min(worksheet.max_row, preview_rows)``.
    ``worksheet.max_row`` is populated from the workbook's optional
    ``<dimension>`` XML tag and is ``None`` when that tag is missing or
    unparseable (verified against openpyxl 3.1.5). ``iter_rows`` does not
    need that value to bound a read — it simply stops after
    ``preview_rows`` rows regardless of whether the declared dimension is
    known, present, or wrong.
    """

    if preview_rows <= 0:
        return []

    preview: list[list[Any]] = []

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=preview_rows,
        values_only=True,
    ):
        preview.append(
            [
                _json_safe_value(value)
                for value in row
            ]
        )

    return preview


# ============================================================================
# Workbook Metadata
# ============================================================================

def _build_workbook_metadata(
    workbook: Any,
) -> ExcelWorkbookMetadata:
    """
    Extract optional workbook properties.
    """

    properties = getattr(
        workbook,
        "properties",
        None,
    )

    if properties is None:
        return ExcelWorkbookMetadata()

    def safe_property(name: str) -> str | None:
        value = getattr(
            properties,
            name,
            None,
        )

        if value is None:
            return None

        if isinstance(
            value,
            (datetime, date, time),
        ):
            return value.isoformat()

        return _safe_string(value)

    return ExcelWorkbookMetadata(
        creator=safe_property("creator"),
        last_modified_by=safe_property(
            "lastModifiedBy"
        ),
        title=safe_property("title"),
        subject=safe_property("subject"),
        description=safe_property(
            "description"
        ),
        category=safe_property("category"),
        keywords=safe_property("keywords"),
        created=safe_property("created"),
        modified=safe_property("modified"),
    )


# ============================================================================
# Hidden Rows / Columns
# ============================================================================
#
# FIX: these three collectors, plus freeze-pane detection below, must be
# called against a NORMAL-mode worksheet object, never against a
# ReadOnlyWorksheet. Verified against openpyxl 3.1.5:
#
#   ReadOnlyWorksheet does not subclass Worksheet at all (its MRO is
#   just (ReadOnlyWorksheet, object)); it only copies over a few
#   methods (iter_rows, cell, values, rows, __getitem__, __iter__).
#   row_dimensions, column_dimensions, merged_cells and freeze_panes
#   are not defined on it, so accessing any of them on a read-only
#   worksheet raises AttributeError immediately.
#
# See "Known openpyxl read-only limitations" in the module docstring
# and scan_excel_workbook()/scan_excel_sheet() below for how the
# normal-mode handle is obtained and passed in.


def _collect_hidden_rows(
    worksheet: Any,
) -> list[ExcelHiddenRowEvidence]:
    """
    Collect explicitly hidden rows.

    Hidden does not mean irrelevant and therefore hidden rows are not removed.

    ``worksheet`` here must be a normal-mode (non-read-only) worksheet.
    """

    hidden: list[ExcelHiddenRowEvidence] = []

    for row_index, dimension in (
        worksheet.row_dimensions.items()
    ):
        if dimension.hidden:
            hidden.append(
                ExcelHiddenRowEvidence(
                    row_index=int(row_index)
                )
            )

    hidden.sort(
        key=lambda item: item.row_index
    )

    return hidden


def _collect_hidden_columns(
    worksheet: Any,
) -> list[ExcelHiddenColumnEvidence]:
    """
    Collect explicitly hidden columns.

    ``worksheet`` here must be a normal-mode (non-read-only) worksheet.
    """

    hidden: list[ExcelHiddenColumnEvidence] = []

    for key, dimension in (
        worksheet.column_dimensions.items()
    ):
        if not dimension.hidden:
            continue

        min_column = getattr(
            dimension,
            "min",
            None,
        )

        max_column = getattr(
            dimension,
            "max",
            None,
        )

        if min_column is None:
            try:
                min_column = (
                    openpyxl.utils.column_index_from_string(
                        str(key)
                    )
                )
            except ValueError:
                continue

        if max_column is None:
            max_column = min_column

        for column_index in range(
            int(min_column),
            int(max_column) + 1,
        ):
            hidden.append(
                ExcelHiddenColumnEvidence(
                    column_index=column_index,
                    column_letter=excel_column_letter(
                        column_index
                    ),
                )
            )

    hidden.sort(
        key=lambda item: item.column_index
    )

    # Remove duplicates while preserving order.
    unique: list[ExcelHiddenColumnEvidence] = []
    seen: set[int] = set()

    for item in hidden:
        if item.column_index in seen:
            continue

        seen.add(item.column_index)
        unique.append(item)

    return unique


# ============================================================================
# Merged Cells
# ============================================================================

def _collect_merged_ranges(
    worksheet: Any,
) -> list[ExcelMergedRangeEvidence]:
    """
    Collect merged ranges in deterministic order.

    ``worksheet`` here must be a normal-mode (non-read-only) worksheet.
    """

    ranges = [
        str(item)
        for item in worksheet.merged_cells.ranges
    ]

    ranges.sort()

    return [
        ExcelMergedRangeEvidence(
            range_reference=item
        )
        for item in ranges
    ]


# ============================================================================
# Metadata Handle (normal-mode, opened solely for sheet-level metadata)
# ============================================================================

def _open_metadata_workbook(
    path: Path,
) -> tuple[Any | None, str | None]:
    """
    Open a second, normal-mode (read_only=False) handle to the same
    workbook, used only to read merged cells, hidden rows/columns and
    freeze panes - none of which openpyxl exposes on a read-only
    worksheet (see module docstring).

    This does load the full workbook into memory, unlike the primary
    read-only handle used for the row-by-row cell scan. For the
    monthly-sized traffic-authority workbooks this project processes
    that cost is acceptable; if it ever becomes a problem for a
    specific source, that is a config-level decision (e.g. skip
    metadata for that source), not something this reader should paper
    over by crashing.

    Returns (workbook, None) on success, or (None, error_message) if
    the normal-mode open itself fails - callers must treat metadata
    collection as unavailable in that case, not as "no metadata found".
    """

    try:
        workbook = openpyxl.load_workbook(
            filename=path,
            read_only=False,
            data_only=False,
            keep_links=True,
        )

        return workbook, None

    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"


# ============================================================================
# Sheet Scan
# ============================================================================

def scan_excel_sheet(
    worksheet: Any,
    sheet_index: int,
    preview_rows: int = PREVIEW_ROWS,
    metadata_worksheet: Any | None = None,
    metadata_error: str | None = None,
) -> ExcelSheetStructuralResult:
    """
    Scan one worksheet in read-only mode.

    The function performs a single structural pass over worksheet rows.

    It does not:
        - detect headers
        - detect totals
        - detect notes
        - detect brands
        - fill blanks
        - normalize values

    Parameters
    ----------
    worksheet:
        The READ-ONLY worksheet used for the row-by-row cell scan.
    metadata_worksheet:
        The matching NORMAL-mode worksheet (same sheet name) used only
        for merged cells / hidden rows & columns / freeze panes. May be
        None if the normal-mode handle could not be opened - see
        metadata_error and ExcelSheetStructuralResult.metadata_read_status.
    metadata_error:
        The error message from opening the normal-mode workbook, if it
        failed. Only used to populate metadata_read_status.
    """

    if preview_rows < 0:
        raise ValueError(
            "preview_rows must be >= 0."
        )

    sheet_name = str(
        worksheet.title
    )

    reported_max_row = int(
        getattr(
            worksheet,
            "max_row",
            0,
        )
        or 0
    )

    reported_max_column = int(
        getattr(
            worksheet,
            "max_column",
            0,
        )
        or 0
    )

    # FIX: worksheet.calculate_dimension() raises ValueError when the
    # sheet's <dimension> XML tag is missing or unparseable (verified
    # against openpyxl 3.1.5: "Worksheet is unsized, use
    # calculate_dimension(force=True)"). reported_dimensions is
    # supplementary evidence only - boundaries (below) already carries
    # the authoritative reported/observed distinction - so on failure
    # we record None rather than losing the whole sheet.
    try:
        reported_dimensions = str(
            worksheet.calculate_dimension()
        )
    except Exception:
        reported_dimensions = None

    preview = _build_preview(
        worksheet,
        preview_rows,
    )

    # FIX: hidden rows/columns, merged ranges and freeze panes are read
    # from the normal-mode metadata_worksheet, never from the read-only
    # `worksheet` (see "Known openpyxl read-only limitations").
    if metadata_worksheet is not None:
        hidden_rows = _collect_hidden_rows(
            metadata_worksheet
        )

        hidden_columns = _collect_hidden_columns(
            metadata_worksheet
        )

        merged_ranges = _collect_merged_ranges(
            metadata_worksheet
        )

        freeze_panes = None

        if metadata_worksheet.freeze_panes is not None:
            freeze_panes = str(
                metadata_worksheet.freeze_panes
            )

        metadata_read_status = "READ_OK"

    else:
        hidden_rows = []
        hidden_columns = []
        merged_ranges = []
        freeze_panes = None

        metadata_read_status = (
            f"UNAVAILABLE: {metadata_error}"
            if metadata_error
            else "UNAVAILABLE: sheet not found in normal-mode workbook"
        )

    rows: list[ExcelRowEvidence] = []

    empty_row_indices: list[int] = []

    observed_first_row: int | None = None
    observed_last_row: int | None = None

    observed_first_column: int | None = None
    observed_last_column: int | None = None

    observed_non_empty_row_count = 0
    observed_non_empty_cell_count = 0

    observed_columns: set[int] = set()

    try:
        for row_cells in worksheet.iter_rows():

            row_number = (
                row_cells[0].row
                if row_cells
                else len(rows) + 1
            )

            non_empty_cells: list[
                ExcelCellEvidence
            ] = []

            non_empty_columns: set[int] = set()

            for cell in row_cells:

                if cell.value is None:
                    continue

                evidence = _build_cell_evidence(
                    cell
                )

                non_empty_cells.append(
                    evidence
                )

                non_empty_columns.add(
                    cell.column
                )

                observed_columns.add(
                    cell.column
                )

                observed_non_empty_cell_count += 1

                if (
                    observed_first_column is None
                    or cell.column
                    < observed_first_column
                ):
                    observed_first_column = (
                        cell.column
                    )

                if (
                    observed_last_column is None
                    or cell.column
                    > observed_last_column
                ):
                    observed_last_column = (
                        cell.column
                    )

            non_empty_count = len(
                non_empty_cells
            )

            is_completely_empty = (
                non_empty_count == 0
            )

            if is_completely_empty:
                empty_row_indices.append(
                    row_number
                )

            else:
                observed_non_empty_row_count += 1

                if observed_first_row is None:
                    observed_first_row = (
                        row_number
                    )

                observed_last_row = (
                    row_number
                )

            first_non_empty_column = (
                min(non_empty_columns)
                if non_empty_columns
                else None
            )

            last_non_empty_column = (
                max(non_empty_columns)
                if non_empty_columns
                else None
            )

            # Only represent empty spans within the row's observed width.
            row_width = max(
                reported_max_column,
                last_non_empty_column or 0,
            )

            empty_spans = _build_empty_spans(
                row_number=row_number,
                non_empty_columns=non_empty_columns,
                max_column=row_width,
            )

            empty_cell_count = sum(
                span.length
                for span in empty_spans
            )

            rows.append(
                ExcelRowEvidence(
                    row_index=row_number,
                    excel_row=row_number,
                    non_empty_count=non_empty_count,
                    empty_cell_count=empty_cell_count,
                    first_non_empty_column=(
                        first_non_empty_column
                    ),
                    last_non_empty_column=(
                        last_non_empty_column
                    ),
                    first_non_empty_column_letter=(
                        excel_column_letter(
                            first_non_empty_column
                        )
                        if first_non_empty_column
                        else None
                    ),
                    last_non_empty_column_letter=(
                        excel_column_letter(
                            last_non_empty_column
                        )
                        if last_non_empty_column
                        else None
                    ),
                    is_completely_empty=(
                        is_completely_empty
                    ),
                    non_empty_cells=tuple(
                        non_empty_cells
                    ),
                    empty_spans=empty_spans,
                )
            )

    except Exception as exc:
        return ExcelSheetStructuralResult(
            sheet_name=sheet_name,
            sheet_index=sheet_index,
            sheet_state=str(
                getattr(
                    worksheet,
                    "sheet_state",
                    "unknown",
                )
            ),
            read_status="READ_ERROR",
            reported_dimensions=(
                reported_dimensions
            ),
            boundaries=None,
            rows=rows,
            empty_row_indices=empty_row_indices,
            empty_column_indices=[],
            empty_column_letters=[],
            merged_ranges=merged_ranges,
            hidden_rows=hidden_rows,
            hidden_columns=hidden_columns,
            metadata_read_status=metadata_read_status,
            freeze_panes=freeze_panes,
            preview=preview,
            preview_rows_requested=preview_rows,
            preview_rows_read=len(preview),
            preview_columns=max(
                (
                    len(row)
                    for row in preview
                ),
                default=0,
            ),
            error=(
                f"{type(exc).__name__}: {exc}"
            ),
        )

    # ------------------------------------------------------------------------
    # Empty columns
    # ------------------------------------------------------------------------

    effective_max_column = max(
        reported_max_column,
        observed_last_column or 0,
    )

    empty_column_indices = [
        column_index
        for column_index in range(
            1,
            effective_max_column + 1,
        )
        if column_index not in observed_columns
    ]

    empty_column_letters = [
        excel_column_letter(
            column_index
        )
        for column_index in empty_column_indices
    ]

    # ------------------------------------------------------------------------
    # Boundaries
    # ------------------------------------------------------------------------

    boundaries = ExcelSheetBoundaries(
        reported_max_row=reported_max_row,
        reported_max_column=reported_max_column,
        observed_first_non_empty_row=(
            observed_first_row
        ),
        observed_last_non_empty_row=(
            observed_last_row
        ),
        observed_first_non_empty_column=(
            observed_first_column
        ),
        observed_last_non_empty_column=(
            observed_last_column
        ),
        observed_first_non_empty_column_letter=(
            excel_column_letter(
                observed_first_column
            )
            if observed_first_column is not None
            else None
        ),
        observed_last_non_empty_column_letter=(
            excel_column_letter(
                observed_last_column
            )
            if observed_last_column is not None
            else None
        ),
        observed_non_empty_row_count=(
            observed_non_empty_row_count
        ),
        observed_non_empty_cell_count=(
            observed_non_empty_cell_count
        ),
    )

    return ExcelSheetStructuralResult(
        sheet_name=sheet_name,
        sheet_index=sheet_index,
        sheet_state=str(
            getattr(
                worksheet,
                "sheet_state",
                "unknown",
            )
        ),
        read_status="READ_OK",
        reported_dimensions=(
            reported_dimensions
        ),
        boundaries=boundaries,
        rows=rows,
        empty_row_indices=empty_row_indices,
        empty_column_indices=empty_column_indices,
        empty_column_letters=empty_column_letters,
        merged_ranges=merged_ranges,
        hidden_rows=hidden_rows,
        hidden_columns=hidden_columns,
        metadata_read_status=metadata_read_status,
        freeze_panes=freeze_panes,
        preview=preview,
        preview_rows_requested=preview_rows,
        preview_rows_read=len(preview),
        preview_columns=max(
            (
                len(row)
                for row in preview
            ),
            default=0,
        ),
        error=None,
    )


# ============================================================================
# Workbook Scan
# ============================================================================

def scan_excel_workbook(
    source: ExcelStructuralInput,
    project_root: Path | None = None,
    preview_rows: int = PREVIEW_ROWS,
) -> ExcelWorkbookStructuralResult:
    """
    Scan one Excel workbook supplied by Discovery.

    The raw workbook is opened read-only and is never modified. A second,
    normal-mode handle is also opened solely to collect worksheet-level
    metadata (merged cells, hidden rows/columns, freeze panes) that
    openpyxl does not expose on read-only worksheets - see the module
    docstring, "Known openpyxl read-only limitations".
    """

    source.validate()

    path = source.resolve_path(
        project_root=project_root
    )

    if not path.exists():
        return ExcelWorkbookStructuralResult(
            file_id=source.file_id,
            file_name=source.file_name,
            relative_path=source.relative_path,
            source_root=source.source_root,
            extension=source.extension.lower(),
            read_status="FILE_NOT_FOUND",
            error=(
                "Source file does not exist: "
                f"{path}"
            ),
        )

    if not path.is_file():
        return ExcelWorkbookStructuralResult(
            file_id=source.file_id,
            file_name=source.file_name,
            relative_path=source.relative_path,
            source_root=source.source_root,
            extension=source.extension.lower(),
            read_status="NOT_A_FILE",
            error=(
                "Resolved source path is not a file: "
                f"{path}"
            ),
        )

    if path.suffix.lower() != ".xlsx":
        return ExcelWorkbookStructuralResult(
            file_id=source.file_id,
            file_name=source.file_name,
            relative_path=source.relative_path,
            source_root=source.source_root,
            extension=source.extension.lower(),
            read_status="UNSUPPORTED_EXTENSION",
            error=(
                "Only .xlsx is supported by the current "
                "Excel Structural reader."
            ),
        )

    workbook = None
    metadata_workbook = None

    try:
        workbook = openpyxl.load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=True,
        )

        metadata_workbook, metadata_error = (
            _open_metadata_workbook(path)
        )

        if metadata_workbook is None:
            LOGGER.warning(
                "Could not open normal-mode metadata handle for %s: %s. "
                "Merged cells, hidden rows/columns and freeze panes will "
                "be marked UNAVAILABLE for every sheet in this file.",
                path,
                metadata_error,
            )

        workbook_metadata = _build_workbook_metadata(
            workbook
        )

        sheets: list[
            ExcelSheetStructuralResult
        ] = []

        sheet_names = list(
            workbook.sheetnames
        )

        for sheet_index, sheet_name in enumerate(
            sheet_names,
            start=1,
        ):
            try:
                worksheet = workbook[sheet_name]

                metadata_worksheet = (
                    metadata_workbook[sheet_name]
                    if metadata_workbook is not None
                    and sheet_name in metadata_workbook.sheetnames
                    else None
                )

                sheet_result = scan_excel_sheet(
                    worksheet=worksheet,
                    sheet_index=sheet_index,
                    preview_rows=preview_rows,
                    metadata_worksheet=metadata_worksheet,
                    metadata_error=metadata_error,
                )

            except Exception as exc:
                sheet_result = ExcelSheetStructuralResult(
                    sheet_name=str(sheet_name),
                    sheet_index=sheet_index,
                    sheet_state="unknown",
                    read_status="READ_ERROR",
                    reported_dimensions=None,
                    boundaries=None,
                    metadata_read_status="UNAVAILABLE: sheet read failed",
                    preview_rows_requested=preview_rows,
                    error=(
                        f"{type(exc).__name__}: {exc}"
                    ),
                )

            sheets.append(sheet_result)

        failed_sheets = [
            sheet
            for sheet in sheets
            if sheet.read_status != "READ_OK"
        ]

        if failed_sheets:
            workbook_status = (
                "PARTIAL_READ_ERROR"
                if len(failed_sheets) < len(sheets)
                else "READ_ERROR"
            )
        else:
            workbook_status = "READ_OK"

        return ExcelWorkbookStructuralResult(
            file_id=source.file_id,
            file_name=source.file_name,
            relative_path=source.relative_path,
            source_root=source.source_root,
            extension=source.extension.lower(),
            read_status=workbook_status,
            reader=READER_NAME,
            reader_version=READER_VERSION,
            reader_schema_version=(
                READER_SCHEMA_VERSION
            ),
            pipeline_version=PIPELINE_VERSION,
            sheet_count=len(sheets),
            sheets=sheets,
            workbook_metadata=workbook_metadata,
            error=None,
        )

    except PermissionError as exc:
        return ExcelWorkbookStructuralResult(
            file_id=source.file_id,
            file_name=source.file_name,
            relative_path=source.relative_path,
            source_root=source.source_root,
            extension=source.extension.lower(),
            read_status="PERMISSION_ERROR",
            error=(
                f"{type(exc).__name__}: {exc}"
            ),
        )

    except Exception as exc:
        return ExcelWorkbookStructuralResult(
            file_id=source.file_id,
            file_name=source.file_name,
            relative_path=source.relative_path,
            source_root=source.source_root,
            extension=source.extension.lower(),
            read_status="READ_ERROR",
            error=(
                f"{type(exc).__name__}: {exc}"
            ),
        )

    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                LOGGER.exception(
                    "Failed to close read-only workbook: %s",
                    path,
                )

        if metadata_workbook is not None:
            try:
                metadata_workbook.close()
            except Exception:
                LOGGER.exception(
                    "Failed to close normal-mode metadata workbook: %s",
                    path,
                )


# ============================================================================
# Discovery Manifest Adapter
# ============================================================================

def excel_input_from_mapping(
    record: dict[str, Any],
) -> ExcelStructuralInput:
    """
    Convert a Discovery manifest record into the minimal Excel input contract.

    This deliberately selects only the fields needed by Excel Structural.

    The full Manifest remains owned by Discovery.
    """

    required_fields = (
        "file_id",
        "relative_path",
        "file_name",
        "extension",
        "source_root",
    )

    missing = [
        field_name
        for field_name in required_fields
        if field_name not in record
    ]

    if missing:
        raise ValueError(
            "Discovery record is missing required fields: "
            + ", ".join(missing)
        )

    return ExcelStructuralInput(
        file_id=str(record["file_id"]),
        relative_path=str(
            record["relative_path"]
        ),
        file_name=str(
            record["file_name"]
        ),
        extension=str(
            record["extension"]
        ).lower(),
        source_root=str(
            record["source_root"]
        ),
        discovery_status=str(
            record.get(
                "discovery_status",
                "DISCOVERED",
            )
        ),
        file_type=str(
            record.get(
                "file_type",
                "excel",
            )
        ),
    )


def excel_inputs_from_manifest_records(
    records: Iterable[dict[str, Any]],
) -> list[ExcelStructuralInput]:
    """
    Convert Discovery records into deterministic Excel input contracts.

    Only:

        discovery_status == DISCOVERED
        file_type == excel
        extension == .xlsx

    are accepted.
    """

    inputs: list[
        ExcelStructuralInput
    ] = []

    for record in records:
        try:
            source = excel_input_from_mapping(
                record
            )
            source.validate()
            inputs.append(source)

        except ValueError as exc:
            LOGGER.warning(
                "Skipping invalid Discovery record: %s",
                exc,
            )

    inputs.sort(
        key=lambda item: (
            item.relative_path.lower(),
            item.file_id,
        )
    )

    return inputs


# ============================================================================
# Multiple Workbook Scan
# ============================================================================

def scan_excel_inputs(
    sources: Iterable[ExcelStructuralInput],
    project_root: Path | None = None,
    preview_rows: int = PREVIEW_ROWS,
) -> list[ExcelWorkbookStructuralResult]:
    """
    Scan multiple Excel files from Discovery input records.

    Files are processed in deterministic order.
    """

    normalized_sources = list(sources)

    normalized_sources.sort(
        key=lambda item: (
            item.relative_path.lower(),
            item.file_id,
        )
    )

    results: list[
        ExcelWorkbookStructuralResult
    ] = []

    for source in normalized_sources:

        result = scan_excel_workbook(
            source=source,
            project_root=project_root,
            preview_rows=preview_rows,
        )

        results.append(result)

    return results


# ============================================================================
# JSON Serialization
# ============================================================================

def workbook_result_to_dict(
    result: ExcelWorkbookStructuralResult,
) -> dict[str, Any]:
    """
    Convert a workbook result into a JSON-safe dictionary.
    """

    return _make_json_serializable(
        result
    )


def save_workbook_result_json(
    result: ExcelWorkbookStructuralResult,
    output_path: Path,
) -> Path:
    """
    Save one workbook's structural evidence as JSON.

    This function never writes to data/raw.
    """

    output_path = Path(
        output_path
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    payload = workbook_result_to_dict(
        result
    )

    temporary_path = output_path.with_suffix(
        output_path.suffix + ".tmp"
    )

    try:
        with temporary_path.open(
            "w",
            encoding="utf-8",
            newline="\n",
        ) as handle:
            json.dump(
                payload,
                handle,
                ensure_ascii=False,
                indent=2,
                sort_keys=False,
            )

        temporary_path.replace(
            output_path
        )

    except Exception:
        try:
            if temporary_path.exists():
                temporary_path.unlink()
        except Exception:
            LOGGER.exception(
                "Failed to remove temporary output: %s",
                temporary_path,
            )

        raise

    return output_path


# ============================================================================
# Stable Output Naming
# ============================================================================

def stable_file_token(
    file_id: str,
) -> str:
    """
    Return a stable filename-safe token.

    The Discovery file_id is normally already hash-like, but hashing it again
    protects the output naming layer from unsafe characters.
    """

    digest = hashlib.sha256(
        file_id.encode(
            "utf-8"
        )
    ).hexdigest()

    return digest[:32]


def structural_output_path(
    result: ExcelWorkbookStructuralResult,
    output_root: Path,
) -> Path:
    """
    Build the recommended output path for one Excel structural result.

    The source month/folder is used only as a physical organization boundary.
    No period is inferred from workbook contents.
    """

    relative = Path(
        result.relative_path
    )

    # The first directory under the configured raw source is used only to
    # mirror source organization. It is not interpreted as a semantic period.
    partition = (
        relative.parts[0]
        if len(relative.parts) > 1
        else "unpartitioned"
    )

    return (
        Path(output_root)
        / "structural"
        / "excel"
        / partition
        / f"{stable_file_token(result.file_id)}.json"
    )


def save_structural_result(
    result: ExcelWorkbookStructuralResult,
    output_root: Path,
) -> Path:
    """
    Save one result using the project's structural output convention.
    """

    output_path = structural_output_path(
        result=result,
        output_root=output_root,
    )

    return save_workbook_result_json(
        result=result,
        output_path=output_path,
    )


# ============================================================================
# Optional Manifest CSV Adapter
# ============================================================================

def load_excel_inputs_from_manifest_csv(
    manifest_path: Path,
) -> list[ExcelStructuralInput]:
    """
    Load Excel inputs directly from the Discovery manifest CSV.

    This is an adapter only.

    It does not perform a raw-directory scan.
    """

    import csv

    manifest_path = Path(
        manifest_path
    )

    if not manifest_path.exists():
        raise FileNotFoundError(
            f"Manifest not found: {manifest_path}"
        )

    records: list[
        dict[str, Any]
    ] = []

    with manifest_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as handle:
        reader = csv.DictReader(
            handle
        )

        for row in reader:
            records.append(
                dict(row)
            )

    return excel_inputs_from_manifest_records(
        records
    )


# ============================================================================
# Validation Helpers
# ============================================================================

def validate_result_traceability(
    result: ExcelWorkbookStructuralResult,
) -> None:
    """
    Validate that the structural result retains source provenance.
    """

    if not result.file_id:
        raise ValueError(
            "Structural result lost file_id."
        )

    if not result.file_name:
        raise ValueError(
            "Structural result lost file_name."
        )

    if not result.relative_path:
        raise ValueError(
            "Structural result lost relative_path."
        )

    if not result.source_root:
        raise ValueError(
            "Structural result lost source_root."
        )

    if result.extension.lower() != ".xlsx":
        raise ValueError(
            "Structural result has an unexpected extension."
        )


# ============================================================================
# Smoke Test
# ============================================================================

def smoke_test_single_file(
    source: ExcelStructuralInput,
    project_root: Path | None = None,
) -> ExcelWorkbookStructuralResult:
    """
    Small programmatic smoke test.

    It scans exactly one Discovery-selected source.

    No directory discovery is performed.
    """

    result = scan_excel_workbook(
        source=source,
        project_root=project_root,
    )

    validate_result_traceability(
        result
    )

    return result


# ============================================================================
# Public API
# ============================================================================

__all__ = [
    "ExcelStructuralInput",
    "ExcelCellEvidence",
    "ExcelEmptyCellSpan",
    "ExcelRowEvidence",
    "ExcelSheetBoundaries",
    "ExcelMergedRangeEvidence",
    "ExcelHiddenRowEvidence",
    "ExcelHiddenColumnEvidence",
    "ExcelWorkbookMetadata",
    "ExcelSheetStructuralResult",
    "ExcelWorkbookStructuralResult",
    "excel_column_letter",
    "classify_excel_value",
    "scan_excel_sheet",
    "scan_excel_workbook",
    "excel_input_from_mapping",
    "excel_inputs_from_manifest_records",
    "scan_excel_inputs",
    "workbook_result_to_dict",
    "save_workbook_result_json",
    "structural_output_path",
    "save_structural_result",
    "load_excel_inputs_from_manifest_csv",
    "validate_result_traceability",
    "smoke_test_single_file",
]


# ============================================================================
# Direct Execution
# ============================================================================

if __name__ == "__main__":
    print(
        "excel_structural.py is a read-only structural reader."
    )
    print(
        "Primary input: ExcelStructuralInput from discovery.py."
    )
    print(
        "Supported Excel format: .xlsx"
    )
    print(
        f"openpyxl version: {READER_VERSION}"
    )